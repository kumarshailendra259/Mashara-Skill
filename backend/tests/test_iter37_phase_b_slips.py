"""iter-37 Phase-B backend tests:
   - Salary Slip Templates: upload/list/delete + RBAC
   - Slip generate / release / download + RBAC + released gate
   - Payroll_run historical employment filter (join/exit windows, prorated working_days)
   - HRMS pagination/search on /api/staff & /api/payroll (X-Total-Count)
   - Staff exit_date persistence
"""
import io
import os
import time
import uuid
import pytest
import requests
from docx import Document

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = {"email": "admin@finance.app", "password": "Admin@123"}


def _login(session, creds):
    r = session.post(f"{API}/auth/login", json=creds, timeout=30)
    assert r.status_code == 200, f"Login failed: {r.status_code} {r.text[:200]}"
    return r


@pytest.fixture(scope="module")
def admin():
    s = requests.Session()
    _login(s, ADMIN)
    return s


@pytest.fixture(scope="module")
def docx_bytes():
    """Build a minimal valid .docx template with a couple of placeholders."""
    d = Document()
    d.add_paragraph("Salary Slip for {{staff_name}} — {{month_name}} {{year}}")
    d.add_paragraph("Net: {{net}} ({{net_words}})")
    d.add_paragraph("Employer: {{company_name}}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------
# Salary Slip Templates — upload / list / delete / RBAC
# ---------------------------------------------------------------
class TestSalarySlipTemplates:
    tpl_id = None

    def test_upload_template_global(self, admin, docx_bytes):
        files = {"file": ("global_slip.docx", docx_bytes,
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = admin.post(f"{API}/salary-slip-templates", files=files, timeout=60)
        assert r.status_code == 201, f"{r.status_code} {r.text[:200]}"
        body = r.json()
        assert body.get("id")
        assert body.get("filename") == "global_slip.docx"
        assert body.get("is_active") is True
        TestSalarySlipTemplates.tpl_id = body["id"]

    def test_upload_deactivates_previous(self, admin, docx_bytes):
        # Upload second global template — first must be deactivated.
        files = {"file": ("global_slip_v2.docx", docx_bytes,
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = admin.post(f"{API}/salary-slip-templates", files=files, timeout=60)
        assert r.status_code == 201
        new_id = r.json()["id"]
        # Older template must now be inactive.
        lst = admin.get(f"{API}/salary-slip-templates", timeout=30).json()
        by_id = {t["id"]: t for t in lst}
        assert by_id[TestSalarySlipTemplates.tpl_id]["is_active"] is False
        assert by_id[new_id]["is_active"] is True
        TestSalarySlipTemplates.tpl_id = new_id  # keep the active one

    def test_list_templates(self, admin):
        r = admin.get(f"{API}/salary-slip-templates", timeout=30)
        assert r.status_code == 200
        assert isinstance(r.json(), list)
        assert any(t["id"] == TestSalarySlipTemplates.tpl_id for t in r.json())

    def test_upload_rejects_non_docx(self, admin):
        files = {"file": ("bad.txt", b"hello", "text/plain")}
        r = admin.post(f"{API}/salary-slip-templates", files=files, timeout=30)
        assert r.status_code == 400

    def test_delete_requires_admin(self, admin):
        # unauth session
        anon = requests.Session()
        r = anon.delete(f"{API}/salary-slip-templates/nope", timeout=30)
        assert r.status_code in (401, 403)


# ---------------------------------------------------------------
# Historical employment filter for payroll_run
# ---------------------------------------------------------------
class TestPayrollEmploymentWindow:
    TAG = f"TEST_ITER37_{uuid.uuid4().hex[:6]}"
    ids: dict = {}

    def _mk(self, admin, name, joining, exit_date=None):
        body = {
            "name": f"{self.TAG}_{name}",
            "designation": "Tester",
            "monthly_salary": 30000,
            "per_day_rate": 0,
            "joining_date": joining,
        }
        if exit_date:
            body["exit_date"] = exit_date
        r = admin.post(f"{API}/staff", json=body, timeout=30)
        assert r.status_code in (200, 201), f"{r.status_code} {r.text[:200]}"
        return r.json()["id"]

    def test_seed_three_staff(self, admin):
        TestPayrollEmploymentWindow.ids["A"] = self._mk(admin, "A", "2026-05-01")
        TestPayrollEmploymentWindow.ids["B"] = self._mk(admin, "B", "2026-06-15")
        TestPayrollEmploymentWindow.ids["C"] = self._mk(
            admin, "C", "2026-05-01", exit_date="2026-07-15"
        )
        # Verify exit_date persisted for C via list endpoint (no per-id GET on /api/staff)
        r = admin.get(f"{API}/staff", params={"q": self.TAG, "limit": 50}, timeout=30)
        assert r.status_code == 200
        c = next((x for x in r.json() if x["id"] == TestPayrollEmploymentWindow.ids["C"]), None)
        assert c is not None, "Staff C not found in list"
        assert c.get("exit_date") == "2026-07-15"

    def _run_and_map(self, admin, month, year):
        r = admin.post(f"{API}/payroll/run", params={"month": month, "year": year}, timeout=120)
        assert r.status_code == 200, f"{r.status_code} {r.text[:200]}"
        # Fetch payroll rows for that month for the 3 staff we care about.
        pr = admin.get(f"{API}/payroll", params={"month": month, "year": year, "limit": 5000}, timeout=60).json()
        watched_ids = set(TestPayrollEmploymentWindow.ids.values())
        return {row["staff_id"]: row for row in pr if row["staff_id"] in watched_ids}

    def test_may_only_a_and_c(self, admin):
        rows = self._run_and_map(admin, 5, 2026)
        assert TestPayrollEmploymentWindow.ids["A"] in rows
        assert TestPayrollEmploymentWindow.ids["C"] in rows
        assert TestPayrollEmploymentWindow.ids["B"] not in rows, "B joined June — must be skipped in May"

    def test_june_all_three_and_b_prorated(self, admin):
        rows = self._run_and_map(admin, 6, 2026)
        for k in ("A", "B", "C"):
            assert TestPayrollEmploymentWindow.ids[k] in rows, f"{k} missing in June"
        b_row = rows[TestPayrollEmploymentWindow.ids["B"]]
        assert b_row["working_days"] == 16, f"B working_days expected 16, got {b_row['working_days']}"
        assert b_row.get("is_partial_month") is True
        # A had a full month
        a_row = rows[TestPayrollEmploymentWindow.ids["A"]]
        assert a_row["working_days"] == 30
        assert a_row.get("is_partial_month") is False

    def test_july_a_and_c_prorated(self, admin):
        rows = self._run_and_map(admin, 7, 2026)
        assert TestPayrollEmploymentWindow.ids["A"] in rows
        assert TestPayrollEmploymentWindow.ids["C"] in rows
        c_row = rows[TestPayrollEmploymentWindow.ids["C"]]
        assert c_row["working_days"] == 15, f"C working_days expected 15, got {c_row['working_days']}"
        assert c_row.get("is_partial_month") is True

    def test_august_only_a(self, admin):
        rows = self._run_and_map(admin, 8, 2026)
        assert TestPayrollEmploymentWindow.ids["A"] in rows
        assert TestPayrollEmploymentWindow.ids["C"] not in rows, "C exited July 15 — must be skipped in Aug"


# ---------------------------------------------------------------
# Slip generate / release / download
# ---------------------------------------------------------------
class TestSlipLifecycle:
    pid = None

    def test_pick_a_payroll_row(self, admin):
        # Pick an A-row we just created for July 2026.
        aid = TestPayrollEmploymentWindow.ids.get("A")
        assert aid, "Staff A must have been created in the earlier test class"
        pr = admin.get(f"{API}/payroll", params={"month": 7, "year": 2026, "limit": 5000}, timeout=60).json()
        row = next((r for r in pr if r["staff_id"] == aid), None)
        assert row, "A payroll row missing for July 2026"
        TestSlipLifecycle.pid = row["id"]

    def test_generate_slip(self, admin):
        r = admin.post(f"{API}/payroll/{TestSlipLifecycle.pid}/generate-slip", timeout=90)
        assert r.status_code == 200, f"{r.status_code} {r.text[:300]}"
        body = r.json()
        assert body.get("ok") is True
        assert body.get("slip_id")
        assert body.get("extension") in ("pdf", "docx")

    def test_download_pre_release_admin_ok(self, admin):
        r = admin.get(f"{API}/payroll/{TestSlipLifecycle.pid}/slip/download", timeout=60)
        assert r.status_code == 200
        assert len(r.content) > 100

    def test_release_slip(self, admin):
        r = admin.patch(f"{API}/payroll/{TestSlipLifecycle.pid}/release-slip",
                        params={"released": True}, timeout=30)
        assert r.status_code == 200
        assert r.json().get("released") is True
        # Verify slip_released_at is set on the payroll doc
        pr = admin.get(f"{API}/payroll", params={"month": 7, "year": 2026, "limit": 5000}, timeout=30).json()
        row = next(x for x in pr if x["id"] == TestSlipLifecycle.pid)
        assert row.get("slip_released_at")

    def test_unrelease_slip(self, admin):
        r = admin.patch(f"{API}/payroll/{TestSlipLifecycle.pid}/release-slip",
                        params={"released": False}, timeout=30)
        assert r.status_code == 200
        assert r.json().get("released") is False


# ---------------------------------------------------------------
# Pagination + search
# ---------------------------------------------------------------
class TestListsPaginationSearch:
    def test_staff_search_and_pagination(self, admin):
        r = admin.get(f"{API}/staff", params={"q": "Boss", "limit": 10}, timeout=30)
        assert r.status_code == 200
        assert "X-Total-Count" in r.headers or "x-total-count" in r.headers
        body = r.json()
        assert isinstance(body, list)
        assert len(body) <= 10
        # Case-insensitive substring match on name/designation/email/etc.
        for s in body:
            hay = " ".join(str(s.get(k) or "") for k in ("name", "designation", "email", "mobile", "employee_code"))
            assert "boss" in hay.lower()

    def test_payroll_search_and_status(self, admin):
        r = admin.get(f"{API}/payroll", params={"limit": 5, "status": "draft"}, timeout=30)
        assert r.status_code == 200
        assert "X-Total-Count" in r.headers or "x-total-count" in r.headers
        body = r.json()
        assert len(body) <= 5
        for row in body:
            assert row.get("status") == "draft"
