"""iter-33 — Offer Letter DOCX fallback when LibreOffice is unavailable.

Production hot-fix validation:
  Code under test:
    - offer_letter.render_offer_letter now returns (bytes, mime_type, extension)
      and falls back to DOCX when `soffice` is missing.
    - server.py:_generate_and_deliver_offer_letter uses tuple return.
    - server.py: /offer-letters/{lid}/download uses stored content_type.
    - server.py: _ensure_libreoffice_installed startup task — idempotent no-op
      when soffice is already installed.

  Scenarios:
    A) Preview HAS LibreOffice → PDF path still works end-to-end (regression).
    B) Force fallback via monkeypatching offer_letter._SOFFICE=None and
       stubbing _find_soffice() → render returns DOCX bytes/mime/ext.
    C) offer_letters DB records carry content_type + extension fields.
    D) Download endpoint sets media_type + Content-Disposition filename to match ext.
    E) _ensure_libreoffice_installed is idempotent — instant return when soffice present.
"""
import io
import os
import sys
import uuid
import asyncio
import importlib
import time

import pytest
import requests
from docx import Document

# Make sure /app/backend is importable so we can pull in `offer_letter` directly
sys.path.insert(0, "/app/backend")

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL",
                          "https://invest-analytics-24.preview.emergentagent.com").rstrip("/")
ADMIN_EMAIL = "admin@finance.app"
ADMIN_PASSWORD = "Admin@123"
RUN = uuid.uuid4().hex[:6]


# ------------------------- helpers -------------------------
def _login(email, pwd):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login",
               json={"email": email, "password": pwd}, timeout=20)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text}"
    return s


def _build_docx(lines):
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


TEMPLATE_LINES = [
    "OFFER LETTER — {{ company_name }}",
    "Date: {{ today }}",
    "Dear {{ staff_name }},",
    "Position: {{ designation }}",
    "Salary: {{ monthly_salary }}",
    "Joining date: {{ joining_date }} at {{ center_name }}.",
    "Login email: {{ login_email }}",
    "Password: {{ login_password }}",
    "Company: {{ company_name }}",
]


# ------------------------- fixtures -------------------------
@pytest.fixture(scope="module")
def admin():
    return _login(ADMIN_EMAIL, ADMIN_PASSWORD)


@pytest.fixture(scope="module")
def created():
    return {"templates": [], "staff": [], "centers": [], "companies": [], "letters": []}


@pytest.fixture(scope="module", autouse=True)
def _cleanup(admin, created):
    yield
    s = admin
    for tid in list(created["templates"]):
        try:
            s.delete(f"{BASE_URL}/api/offer-letter-templates/{tid}", timeout=15)
        except Exception:
            pass
    if created["staff"]:
        try:
            s.post(f"{BASE_URL}/api/entities/staff/bulk-delete",
                   json={"ids": created["staff"]}, timeout=30)
        except Exception:
            pass
    if created["centers"]:
        try:
            s.post(f"{BASE_URL}/api/entities/center/bulk-delete",
                   json={"ids": created["centers"]}, timeout=30)
        except Exception:
            pass
    if created["companies"]:
        try:
            s.post(f"{BASE_URL}/api/entities/company/bulk-delete",
                   json={"ids": created["companies"]}, timeout=30)
        except Exception:
            pass


@pytest.fixture(scope="module")
def seed(admin, created):
    """Company + center linked, plus an active per-company template."""
    s = admin
    # company
    r = s.post(f"{BASE_URL}/api/entities/company", json={
        "name": f"TEST_iter33_Co_{RUN}",
        "address": "1 Test Rd", "city": "Bengaluru", "state": "KA",
        "gst_number": "29TEST33GST01ZK", "pan_number": "TEST33PAN9",
        "email": "hr@testco33.local", "mobile": "9990000033",
    }, timeout=20)
    assert r.status_code in (200, 201), r.text
    comp = r.json()
    created["companies"].append(comp["id"])

    # center bound to company
    r = s.post(f"{BASE_URL}/api/entities/center", json={
        "name": f"TEST_iter33_Center_{RUN}", "city": "Bengaluru", "state": "KA",
    }, timeout=20)
    assert r.status_code in (200, 201), r.text
    center = r.json()
    created["centers"].append(center["id"])
    r = s.put(f"{BASE_URL}/api/entities/center/{center['id']}", json={
        "name": center["name"], "city": "Bengaluru", "state": "KA",
        "company_id": comp["id"],
    }, timeout=20)
    assert r.status_code in (200, 201), r.text

    # per-company template
    docx = _build_docx(TEMPLATE_LINES)
    r = s.post(f"{BASE_URL}/api/offer-letter-templates",
               files={"file": (f"TEST_iter33_tpl_{RUN}.docx", docx,
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
               params={"company_id": comp["id"]}, timeout=30)
    assert r.status_code == 201, r.text
    tpl_id = r.json()["id"]
    created["templates"].append(tpl_id)

    return {
        "company_id": comp["id"], "company_name": comp["name"],
        "center_id": center["id"], "center_name": center["name"],
        "template_id": tpl_id,
    }


# =========================================================================
# A) PDF happy path — preview environment has soffice
# =========================================================================
class TestPdfHappyPath:
    """With LibreOffice available, offer letter must be delivered as PDF."""

    def test_create_staff_generates_pdf(self, admin, seed, created):
        staff_email = f"iter33_pdf_{RUN}@finance.local"
        r = admin.post(f"{BASE_URL}/api/staff", json={
            "name": f"TEST_iter33_PdfStaff_{RUN}",
            "designation": "Senior Trainer",
            "monthly_salary": 45000,
            "per_day_rate": 1500,
            "joining_date": "2026-02-01",
            "center_id": seed["center_id"],
            "email": staff_email,
            "mobile": "9998887701",
            "create_login": True,
            "send_credentials_email": False,
            "send_offer_letter": True,
        }, timeout=180)
        assert r.status_code in (200, 201), f"{r.status_code} {r.text}"
        data = r.json()
        created["staff"].append(data["id"])

        ols = data.get("offer_letter_status") or {}
        assert ols.get("generated") is True, f"not generated: {ols}"
        assert ols.get("letter_id"), f"no letter_id: {ols}"
        assert ols.get("extension") == "pdf", \
            f"expected extension='pdf', got '{ols.get('extension')}' — LibreOffice may be missing on preview! full ols={ols}"
        assert ols.get("pdf_path", "").endswith(".pdf"), \
            f"pdf_path should end .pdf, got {ols.get('pdf_path')}"
        created["_pdf_letter_id"] = ols["letter_id"]

    def test_offer_letters_record_carries_content_type_and_extension(self, admin, created):
        """The offer_letters MongoDB doc (via /offer-letters listing) must include
        content_type + extension fields alongside pdf_path."""
        lid = created.get("_pdf_letter_id")
        assert lid
        rows = admin.get(f"{BASE_URL}/api/offer-letters", timeout=15).json()
        found = next((x for x in rows if x["id"] == lid), None)
        assert found, f"letter {lid} not in listing"
        assert found.get("content_type") == "application/pdf", \
            f"missing/wrong content_type: {found.get('content_type')}"
        assert found.get("extension") == "pdf", \
            f"missing/wrong extension: {found.get('extension')}"
        assert found.get("pdf_path", "").endswith(".pdf")
        assert found.get("filename", "").endswith(".pdf"), \
            f"filename should be OfferLetter_Name.pdf, got {found.get('filename')}"

    def test_download_returns_pdf_with_matching_disposition(self, admin, created):
        lid = created.get("_pdf_letter_id")
        assert lid
        r = admin.get(f"{BASE_URL}/api/offer-letters/{lid}/download", timeout=30)
        assert r.status_code == 200, f"{r.status_code} {r.text[:400]}"
        assert r.headers.get("content-type", "").startswith("application/pdf"), \
            f"content-type: {r.headers.get('content-type')}"
        cd = r.headers.get("content-disposition", "")
        assert ".pdf" in cd, f"Content-Disposition filename must end .pdf: {cd}"
        assert r.content[:5] == b"%PDF-", \
            f"bytes are not a PDF: {r.content[:16]!r}"


# =========================================================================
# B) Force DOCX fallback path — monkeypatch offer_letter._SOFFICE=None
# =========================================================================
class TestDocxFallback:
    """Directly test render_offer_letter() with LibreOffice hidden.
    This exercises the exact same code path that server.py uses to render."""

    def test_render_returns_docx_when_soffice_missing(self, monkeypatch):
        import offer_letter
        importlib.reload(offer_letter)  # fresh module state

        monkeypatch.setattr(offer_letter, "_SOFFICE", None)
        monkeypatch.setattr(offer_letter, "_find_soffice", lambda: None)

        template_bytes = _build_docx(TEMPLATE_LINES)
        staff = {
            "name": "Alice Fallback", "designation": "QA Engineer",
            "monthly_salary": 30000, "joining_date": "2026-03-01",
            "email": "alice@fallback.test",
        }
        company = {"name": "Fallback Co", "address": "1 Street",
                   "city": "BLR", "state": "KA", "gst_number": "GST", "pan_number": "PAN"}
        center = {"name": "Fallback Center"}
        letter_bytes, ct, ext = offer_letter.render_offer_letter(
            template_bytes=template_bytes, staff=staff, company=company,
            center=center, login_email="alice@fallback.test", login_password="TempPw@1234",
        )
        assert isinstance(letter_bytes, (bytes, bytearray))
        assert ext == "docx", f"expected docx fallback ext, got {ext!r}"
        assert ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
            f"wrong mime for docx fallback: {ct}"
        # DOCX is a zip → begins with PK
        assert letter_bytes[:2] == b"PK", \
            f"DOCX bytes should start with PK zip magic, got {letter_bytes[:8]!r}"

        # Confirm the DOCX we got actually rendered — reopen it and look for staff name
        doc = Document(io.BytesIO(letter_bytes))
        text_all = "\n".join(p.text for p in doc.paragraphs)
        assert "Alice Fallback" in text_all, f"staff_name missing from DOCX text:\n{text_all}"
        assert "QA Engineer" in text_all, "designation missing from DOCX"
        assert "Fallback Co" in text_all, "company_name missing from DOCX"
        # No un-substituted Jinja placeholders
        assert "{{" not in text_all, f"raw Jinja placeholder left in DOCX:\n{text_all}"

    def test_docx_to_pdf_returns_none_when_soffice_missing(self, monkeypatch, tmp_path):
        import offer_letter
        importlib.reload(offer_letter)

        monkeypatch.setattr(offer_letter, "_SOFFICE", None)
        monkeypatch.setattr(offer_letter, "_find_soffice", lambda: None)

        # Create a real docx on disk
        docx_bytes = _build_docx(["dummy"])
        src = tmp_path / "x.docx"
        src.write_bytes(docx_bytes)
        result = offer_letter._docx_to_pdf(str(src), str(tmp_path))
        assert result is None, \
            f"_docx_to_pdf should return None when soffice missing, got {result!r}"


# =========================================================================
# C) _ensure_libreoffice_installed idempotency
# =========================================================================
class TestEnsureLibreofficeIdempotent:
    """When soffice already exists, the startup task must return immediately."""

    def test_function_exists_and_signature(self):
        # Import server without triggering FastAPI startup (just grab the coroutine func)
        import server
        assert hasattr(server, "_ensure_libreoffice_installed"), \
            "_ensure_libreoffice_installed missing from server"
        fn = server._ensure_libreoffice_installed
        assert asyncio.iscoroutinefunction(fn), "must be async"
        # Should take zero positional args (just self-contained bg task)
        import inspect
        sig = inspect.signature(fn)
        assert len(sig.parameters) == 0, \
            f"expected no params, got {list(sig.parameters)}"

    def test_fast_return_when_soffice_present(self):
        """Preview has soffice installed, so calling directly should be a no-op
        that returns near-instantly (well under 5s — no apt-get invoked)."""
        import server
        loop = asyncio.new_event_loop()
        try:
            t0 = time.perf_counter()
            loop.run_until_complete(server._ensure_libreoffice_installed())
            elapsed = time.perf_counter() - t0
        finally:
            loop.close()
        # A real apt-get install takes 30–120s. Idempotent no-op should be <1s.
        assert elapsed < 5.0, \
            f"_ensure_libreoffice_installed took {elapsed:.2f}s — expected <1s no-op when soffice present"

    def test_no_op_never_raises_even_without_apt(self, monkeypatch):
        """If apt-get is not available AND soffice is not present, must silently no-op."""
        import server
        import shutil as sh

        real_which = sh.which

        def fake_which(cmd):
            if cmd in ("soffice", "libreoffice", "apt-get"):
                return None
            return real_which(cmd)

        # server._ensure_libreoffice_installed imports shutil inside the function
        # (as `import shutil as _shutil`), so patch the module-level shutil.which
        monkeypatch.setattr(sh, "which", fake_which)

        loop = asyncio.new_event_loop()
        try:
            # Should NOT raise, and return quickly
            loop.run_until_complete(server._ensure_libreoffice_installed())
        finally:
            loop.close()


# =========================================================================
# D) Regression sanity — offer_letter._find_soffice detects at call time
# =========================================================================
class TestFindSofficeAtCallTime:
    def test_find_soffice_re_detects(self):
        """The fix says: _find_soffice() is called on every render attempt so a
        later apt install becomes visible without restarting the app."""
        import offer_letter
        importlib.reload(offer_letter)
        # In preview, soffice IS installed
        detected = offer_letter._find_soffice()
        assert detected, "soffice should be detected in preview environment"
        assert "soffice" in detected or "libreoffice" in detected, \
            f"unexpected path: {detected}"
